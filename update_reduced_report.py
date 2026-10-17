"""
update_reduced_report.py
Refines CPPE_Mini_project_2_report_KEISHORE 345678.docx to:
1. Retain only the 8 most essential, high-impact screenshots (2 per module) without redundant repeats.
2. Embed the 5 publication-quality architecture & activity diagrams.
3. Update all technical text and schema data to match the unified SmartShelf platform.
4. Remove redundant screenshot tables and bullet paragraphs cleanly.
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
REPORT_PATH = os.path.join(PROJECT_DIR, "CPPE_Mini_project_2_report_KEISHORE 345678.docx")
BACKUP_PATH = os.path.join(PROJECT_DIR, "CPPE_Mini_project_2_report_KEISHORE 345678_BACKUP.docx")

DIAGRAMS_DIR = os.path.join(PROJECT_DIR, "diagrams")
STITCH_DIR = os.path.join(PROJECT_DIR, "stitch_smartshelf_enterprise_inventory_suite")
SCREENSHOTS_DIR = os.path.join(PROJECT_DIR, "screenshots")

doc = docx.Document(BACKUP_PATH)
print("Loaded document from backup. Total paragraphs:", len(doc.paragraphs), "Total tables:", len(doc.tables))

# -------------------------------------------------------------
# 1. UPDATE PARAGRAPH TEXT DETAILS (Matching Project Architecture)
# -------------------------------------------------------------
print("\n--- Updating Paragraph Text to Match SmartShelf Architecture ---")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    
    # Abstract update
    if "The application is developed using Python and Streamlit, with SQLite used to store product information." in txt:
        p.text = (
            "The application is developed using Python with FastAPI for the backend REST engine, "
            "Google Stitch UI Suite for the primary enterprise single-page application, and SQLite for "
            "ACID-compliant transactional data storage with WAL mode, while retaining Streamlit for "
            "interactive data analytics. Data processing is handled with Pandas and deterministic business logic."
        )
        print(f"Updated Abstract P{i}")
        
    elif "Keywords: Inventory Management, Expiry Tracking, Stock Monitoring, Python, Streamlit, SQLite." in txt:
        p.text = "Keywords: Inventory Management, Expiry Tracking, Stock Monitoring, Python, FastAPI, Google Stitch UI, Streamlit, SQLite."
        print(f"Updated Keywords P{i}")
        
    elif "SmartShelf is built with Python for application logic, Streamlit for the web interface, SQLite for data storage" in txt:
        p.text = (
            "SmartShelf is built with Python and FastAPI for application and REST service logic, "
            "Google Stitch UI Suite for the enterprise web interface, SQLite for transactional data storage, "
            "Pandas for data processing, and Streamlit as an integrated secondary analytics studio."
        )
        print(f"Updated Scope P{i}")
        
    elif txt.startswith("Presentation Layer: Streamlit provides the user interface"):
        p.text = (
            "Presentation Layer: Google Stitch Enterprise UI provides the modern single-page SaaS interface "
            "for inventory tracking, catalog inspection, product registration, lifecycle edits, exception dispatch, "
            "and reporting. In addition, an embedded Streamlit Analytics Studio provides native multi-page data exploration."
        )
        print(f"Updated Presentation Layer P{i}")
        
    elif txt.startswith("Application Layer: Python processes user inputs"):
        p.text = (
            "Application Layer: Python with FastAPI (ASGI Uvicorn) acts as the high-throughput asynchronous backend service. "
            "It exposes REST API endpoints (/api/products, /api/alerts/summary, /api/audit-logs) with automatic OpenAPI "
            "Swagger documentation, applies strict validation, and evaluates deterministic expiry and stock rules."
        )
        print(f"Updated Application Layer P{i}")
        
    elif txt.startswith("Database Layer: SQLite stores product details"):
        p.text = (
            "Database Layer: SQLite 3 stores product details and audit trails in a persistent, transactional database "
            "(smartshelf.db) configured with Write-Ahead Logging (WAL) for high concurrency and zero server maintenance."
        )
        print(f"Updated Database Layer P{i}")
        
    elif txt == "User → Streamlit Interface → Python Processing → SQLite Database → Inventory Analysis and Alerts → Results displayed to User.":
        p.text = "User → Google Stitch UI / Browser → FastAPI REST API → SQLite Database (WAL) → Expiry & Stock Rule Engine → Dynamic Updates displayed in UI."
        print(f"Updated Workflow P{i}")
        
    elif txt.startswith("When the user submits a request such as adding a product, the Streamlit interface passes"):
        p.text = (
            "When the user submits a request such as adding a product, the Stitch UI passes the data via asynchronous "
            "HTTP JSON fetch to the FastAPI layer, which validates the payload against Pydantic models and stores it in "
            "the SQLite database using parameterised SQL. When the inventory is displayed, stored records are retrieved, "
            "analysed against the expiry and stock rules, and the resulting JSON dataset is rendered dynamically in the browser."
        )
        print(f"Updated Request Flow P{i}")
        
    elif txt == "Develop the application using Python and Streamlit.":
        p.text = "Develop the unified application using Python, FastAPI, and Google Stitch UI, retaining Streamlit for supplementary analytics."
        print(f"Updated Dev Step P{i}")
        
    elif txt == "streamlit run app.py":
        p.text = "python main.py 5000    # (FastAPI + Stitch UI)  |  streamlit run app.py (Analytics Studio)"
        print(f"Updated Command P{i}")
        
    elif txt.startswith("The terminal shows the Streamlit server started"):
        p.text = "The terminal shows the unified FastAPI Uvicorn server started on port 5000 and the local URLs for web and API access."
        print(f"Updated Terminal Desc P{i}")
        
    elif txt.startswith("The Python environment is configured with Streamlit and Pandas."):
        p.text = "The Python environment is configured with FastAPI, Uvicorn, Streamlit, and Pandas."
        print(f"Updated Env Desc P{i}")
        
    elif txt == "The app is run locally with streamlit run, or pushed to GitHub and deployed to the cloud.":
        p.text = "The app is run locally with python main.py 5000 (or run_enterprise.bat), or pushed to GitHub and deployed to the cloud."
        print(f"Updated Run Desc P{i}")

# -------------------------------------------------------------
# 2. UPDATE SCREENSHOT DESCRIPTIVE BULLET POINTS
# -------------------------------------------------------------
print("\n--- Updating Screenshot Section Bullets for 2 Essential Screenshots Per Module ---")

# Module 1 (Data Engineering)
# P165 & P166 update, P167 & P168 remove
for i, p in enumerate(doc.paragraphs):
    if "The home page shows the application title and navigation" in p.text:
        p.text = "The product registration form collects trade name, category, quantity, unit price, and expiry date with real-time valuation and live preview."
    elif "The form collects the product name, category, quantity" in p.text:
        p.text = "The inventory catalog inspector lists all stored products with keyword search, category filtering, stock status badges, and item profile inspector drawer."
    elif "All stored products are listed in a sortable table with the total number" in p.text:
        p.text = "" # cleared
    elif "A product is selected to modify its details or to remove it permanently" in p.text:
        p.text = "" # cleared

# Module 2 (Cloud Engineering)
# P205 & P206 update, P207 & P208 remove
for i, p in enumerate(doc.paragraphs):
    if "The terminal shows the unified FastAPI Uvicorn server" in p.text:
        p.text = "The terminal shows the unified FastAPI Uvicorn server started on port 5000, database initialized with WAL mode, and active REST routes."
    elif "The deployment settings show the GitHub repository" in p.text:
        p.text = "The OpenAPI Swagger documentation demonstrates the interactive REST API endpoints and cloud deployment schema at /docs."
    elif "SmartShelf is opened in a browser through its public web address" in p.text:
        p.text = "" # cleared
    elif "If only a local prototype is developed, include screenshots" in p.text:
        p.text = "" # cleared

# Module 3 (AI Engineering)
# P257 & P258 update, P259 & P260 remove
for i, p in enumerate(doc.paragraphs):
    if "Metrics show the total products, expired products, near-expiry" in p.text:
        p.text = "The Alerts and Expiration Exception Hub highlights expired batches, imminent expiration warnings (<= 7 days), and low-stock replenishment items."
    elif "A table lists products that will expire within the configured" in p.text:
        p.text = "The Reports and Data Export Center summarises inventory stock distribution across categories and facilitates CSV/JSON data auditing."
    elif "Products whose quantity is below the minimum level are highlighted" in p.text:
        p.text = "" # cleared
    elif "A chart or report summarises stock quantities by category" in p.text:
        p.text = "" # cleared

# Module 4 (Security Engineering)
# P297 & P298 update, P299 & P300 remove
for i, p in enumerate(doc.paragraphs):
    if "An error message appears when a negative or non-numeric quantity" in p.text:
        p.text = "An input validation error notification appears when an invalid or negative quantity is entered, safely rejecting invalid submissions."
    elif "The form refuses to save when required fields are left empty" in p.text:
        p.text = "The system audit log provides an immutable governance ledger tracking all user actions (ADD, EDIT, DELETE) with complete timestamps and parameterised query protection."
    elif "Only authorized users can reach the add, edit, and delete functions" in p.text:
        p.text = "" # cleared
    elif "A confirmation message is shown after a valid record is saved" in p.text:
        p.text = "" # cleared

# Clean up empty paragraphs that were cleared above
paras_to_remove = []
for p in doc.paragraphs:
    # If it was one of the cleared bullets (originally had empty text now)
    # We only remove paragraphs if they are empty AND part of the bullet lists
    pass # Leaving clean 0-height empty runs or removing element:
# Let's remove elements of empty paragraphs that have no text and were previously screenshot bullets
for p in list(doc.paragraphs):
    if p.text == "" and p.paragraph_format.left_indent is not None:
        p._element.getparent().remove(p._element)

# -------------------------------------------------------------
# 3. UPDATE DATA TABLES (Table 1 & Table 3)
# -------------------------------------------------------------
print("\n--- Updating Data Tables ---")

# Table 1: Technology Stack by Layer
tbl1 = doc.tables[1]
if len(tbl1.rows) >= 4:
    tbl1.rows[1].cells[0].text = "Presentation"
    tbl1.rows[1].cells[1].text = "Google Stitch UI Suite & Streamlit"
    tbl1.rows[1].cells[2].text = "Enterprise responsive dashboard, catalog inspector, registration, alerts, audit log, and embedded analytics studio"
    
    tbl1.rows[2].cells[0].text = "Application"
    tbl1.rows[2].cells[1].text = "Python, FastAPI, Pandas"
    tbl1.rows[2].cells[2].text = "Asynchronous REST API, Swagger docs, parameterised CRUD, rule evaluation, and audit governance"
    
    tbl1.rows[3].cells[0].text = "Database"
    tbl1.rows[3].cells[1].text = "SQLite 3 (WAL Mode)"
    tbl1.rows[3].cells[2].text = "Persistent transactional storage in smartshelf.db with complete audit trails"
    print("Updated Table 1 (Technology Stack)")

# Table 3: Sample Product Records
tbl3 = doc.tables[3]
sample_data = [
    ("1", "Organic Whole Milk 1L", "Dairy", "40", "28.00", "2026-10-20"),
    ("2", "Artisan French Baguette", "Bakery", "15", "55.00", "2026-10-18"),
    ("3", "Fresh Atlantic Salmon Fillet", "Seafood", "6", "120.00", "2026-10-14"),
    ("4", "Cold Brew Bottled Coffee", "Beverages", "25", "95.00", "2026-12-05"),
]
for r_idx, row_values in enumerate(sample_data):
    if r_idx + 1 < len(tbl3.rows):
        for c_idx, val in enumerate(row_values):
            tbl3.rows[r_idx + 1].cells[c_idx].text = val
print("Updated Table 3 (Sample Product Records to actual seed data)")

# -------------------------------------------------------------
# 4. INSERT THE 5 ARCHITECTURE & ACTIVITY DIAGRAMS
# -------------------------------------------------------------
print("\n--- Inserting 5 Publication-Quality Diagrams ---")

diagram_map = {
    "Figure 3.1: SmartShelf System Architecture": os.path.join(DIAGRAMS_DIR, "fig_3_1_architecture.png"),
    "Figure 4.1: Activity Diagram for Data Engineering": os.path.join(DIAGRAMS_DIR, "fig_4_1_data_engineering.png"),
    "Figure 4.2: Activity Diagram for Cloud Engineering": os.path.join(DIAGRAMS_DIR, "fig_4_2_cloud_engineering.png"),
    "Figure 4.3: Activity Diagram for AI Engineering": os.path.join(DIAGRAMS_DIR, "fig_4_3_ai_engineering.png"),
    "Figure 4.4: Activity Diagram for Security Engineering": os.path.join(DIAGRAMS_DIR, "fig_4_4_security_engineering.png"),
}

for fig_caption, img_path in diagram_map.items():
    found = False
    for i, p in enumerate(doc.paragraphs):
        if fig_caption in p.text:
            found = True
            target_p = doc.paragraphs[i - 1]
            target_p.text = ""
            target_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            target_p.paragraph_format.space_before = Pt(8)
            target_p.paragraph_format.space_after = Pt(4)
            run = target_p.add_run()
            run.add_picture(img_path, width=Inches(6.0))
            
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(12)
            print(f"Inserted diagram '{os.path.basename(img_path)}' above caption at P{i}")
            break
    if not found:
        print(f"WARNING: Caption not found: {fig_caption}")

# -------------------------------------------------------------
# 5. POPULATE THE 8 ESSENTIAL SCREENSHOTS IN ORIGINAL TABLES
# -------------------------------------------------------------
print("\n--- Populating the 8 Essential Screenshots ---")

# In the original document (before deletion):
# Table 5 -> Module 1, Screenshot 1 (Add Product Form)
# Table 6 -> Module 1, Screenshot 2 (Catalog Inspector)
# Table 9 -> Module 2, Screenshot 1 (Terminal Server)
# Table 10 -> Module 2, Screenshot 2 (Swagger OpenAPI Docs)
# Table 15 -> Module 3, Screenshot 1 (Alerts & Expiration Hub)
# Table 17 -> Module 3, Screenshot 2 (Reports & Data Export)
# Table 20 -> Module 4, Screenshot 1 (Validation Error Quantity)
# Table 22 -> Module 4, Screenshot 2 (System Audit Log)

essential_screenshots = {
    5: {
        "img": os.path.join(STITCH_DIR, "register_new_product_form", "screen.png"),
        "caption": "Screenshot 4.1.1: Register New Product Form with Real-Time Valuation & Preview",
    },
    6: {
        "img": os.path.join(STITCH_DIR, "inventory_catalog_inspector", "screen.png"),
        "caption": "Screenshot 4.1.2: Inventory Catalog Inspector with Filtering and Item Profile Drawer",
    },
    9: {
        "img": os.path.join(SCREENSHOTS_DIR, "terminal_server_running.png"),
        "caption": "Screenshot 4.2.1: Unified Application Server Running on Port 5000 with FastAPI & SQLite",
    },
    10: {
        "img": os.path.join(SCREENSHOTS_DIR, "swagger_api_docs.png"),
        "caption": "Screenshot 4.2.2: OpenAPI Swagger REST Endpoints and Server Configuration",
    },
    15: {
        "img": os.path.join(STITCH_DIR, "alerts_expiration_exception_hub", "screen.png"),
        "caption": "Screenshot 4.3.1: Alerts & Expiration Exception Hub Displaying Imminent Expiry and Low-Stock Batches",
    },
    17: {
        "img": os.path.join(STITCH_DIR, "reports_data_export_center", "screen.png"),
        "caption": "Screenshot 4.3.2: Category Stock Distribution and Data Export Center",
    },
    20: {
        "img": os.path.join(SCREENSHOTS_DIR, "validation_error_quantity.png"),
        "caption": "Screenshot 4.4.1: Input Validation Error Rejection for Negative Quantity",
    },
    22: {
        "img": os.path.join(STITCH_DIR, "system_audit_log_governance_ledger", "screen.png"),
        "caption": "Screenshot 4.4.2: System Audit Log & Governance Ledger Tracking Secure Operations",
    },
}

for tbl_idx, config in essential_screenshots.items():
    tbl = doc.tables[tbl_idx]
    cell = tbl.cell(0, 0)
    img_path = config["img"]
    caption_text = config["caption"]
    
    p_img = cell.paragraphs[0]
    p_img.text = ""
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(5.7))
    
    if len(cell.paragraphs) > 1:
        p_cap = cell.paragraphs[1]
    else:
        p_cap = cell.add_paragraph()
        
    p_cap.text = ""
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(6)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Arial"
    run_cap.font.size = Pt(10)
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(30, 41, 59)
    print(f"Populated essential screenshot in Table {tbl_idx}: {caption_text}")

# -------------------------------------------------------------
# 6. REMOVE REDUNDANT SCREENSHOT TABLES (IN DESCENDING ORDER)
# -------------------------------------------------------------
print("\n--- Removing Redundant Tables in Descending Order ---")
redundant_table_indices = [23, 21, 16, 14, 11, 7, 4]

for idx in redundant_table_indices:
    tbl = doc.tables[idx]
    tbl._element.getparent().remove(tbl._element)
    print(f"Removed redundant Table {idx}")

print("Total tables remaining after reduction:", len(doc.tables))

# -------------------------------------------------------------
# 7. SAVE REFINED REPORT
# -------------------------------------------------------------
doc.save(REPORT_PATH)
file_size = os.path.getsize(REPORT_PATH)
print(f"\nSUCCESS! Refined report saved to:\n{REPORT_PATH}\nFile size: {file_size / (1024*1024):.2f} MB")
